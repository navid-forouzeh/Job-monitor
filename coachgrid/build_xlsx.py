#!/usr/bin/env python3
"""
Baut aus den CSVs in coachgrid/data/ eine Excel-Datei.

    python3 coachgrid/build_xlsx.py   -> coachgrid/Vereine_Schweiz.xlsx

Drei Blätter: Übersicht (Zahlen pro Sportart), Alle Vereine (die Liste mit
Autofilter), Quellen.
"""

import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(BASE, "data")
OUT = os.path.join(BASE, "Vereine_Schweiz.xlsx")
STAND = "28.07.2026"

# Datei -> (Sportart, Priorität gemäss marktanalyse-vereine.md)
SOURCES = [
    ("golf.csv", "Golf", "A – eigene Anlage, angestelltes Personal"),
    ("tennis.csv", "Tennis", "B – eigene Anlage"),
    ("schwimmen.csv", "Schwimmen", "B – Kursgeschäft"),
    ("reiten_eishockey.csv", "Reiten & Eishockey", "A/B – meist kommerziell geführt"),
]

# HTTP-Code aus dem Link-Check -> Klartext + Farbe
STATUS = {
    "200": ("Link ok", "C6E6D0"),
    "301": ("Link ok (Weiterleitung)", "C6E6D0"),
    "403": ("Prüfen – Seite blockt Bots", "FBE7C0"),
    "406": ("Prüfen – Seite blockt Bots", "FBE7C0"),
    "429": ("Prüfen – Seite blockt Bots", "FBE7C0"),
    "000": ("Prüfen – Zeitüberschreitung", "FBE7C0"),
    "334": ("Prüfen – unklare Antwort", "FBE7C0"),
    "404": ("Link tot (404)", "F5CDC8"),
    "500": ("Serverfehler (500)", "F5CDC8"),
    "-": ("keine Website hinterlegt", "EFEFEF"),
    "": ("keine Website hinterlegt", "EFEFEF"),
}

HEADERS = [
    ("Sportart", 20),
    ("Name", 46),
    ("Telefon", 20),
    ("E-Mail", 34),
    ("Website", 40),
    ("Link geprüft", 26),
    ("PLZ", 8),
    ("Ort", 24),
    ("Strasse", 28),
    ("Angebot", 46),
    ("Karte", 10),
    ("Priorität", 34),
]

ARIAL = "Arial"
INK = "1F3B33"
HEAD_FILL = PatternFill("solid", fgColor=INK)
LINK_FONT = Font(name=ARIAL, size=10, color="0F6F5C", underline="single")
BODY = Font(name=ARIAL, size=10)


def load():
    rows = []
    for filename, sport, prio in SOURCES:
        with open(os.path.join(DATA, filename), encoding="utf-8") as handle:
            for row in csv.DictReader(handle, delimiter=";"):
                if not row.get("name"):
                    continue
                coords = (row.get("koordinaten") or "").strip()
                ort = (row.get("ort") or "").strip()
                if coords:
                    lat, lon = coords.split(",")
                    maps = f"https://www.openstreetmap.org/?mlat={lat}&mlon={lon}&zoom=16"
                elif ort:
                    query = (row["name"] + " " + ort).replace(" ", "+")
                    maps = f"https://www.openstreetmap.org/search?query={query}"
                else:
                    maps = ""
                rows.append(
                    {
                        "sport": sport,
                        "prio": prio,
                        "name": row["name"].strip(),
                        "tel": (row.get("telefon") or "").strip(),
                        "mail": (row.get("email") or "").strip(),
                        "web": (row.get("website") or "").strip(),
                        "status": (row.get("link_status") or "").strip(),
                        "plz": (row.get("plz") or "").strip(),
                        "ort": ort,
                        "strasse": (row.get("strasse") or "").strip(),
                        "angebot": (row.get("angebot") or "").strip(),
                        "maps": maps,
                    }
                )
    rows.sort(key=lambda r: (r["sport"], r["name"]))
    return rows


def write_list(sheet, rows):
    sheet.freeze_panes = "B2"
    for index, (title, width) in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 22

    for offset, row in enumerate(rows):
        line = offset + 2
        label, colour = STATUS.get(row["status"], (f"HTTP {row['status']}", "FBE7C0"))
        values = [
            row["sport"],
            row["name"],
            row["tel"] or None,
            row["mail"] or None,
            row["web"] or None,
            label,
            row["plz"] or None,
            row["ort"] or None,
            row["strasse"] or None,
            row["angebot"] or None,
            "Karte" if row["maps"] else None,
            row["prio"],
        ]
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=line, column=index, value=value)
            cell.font = BODY
            cell.alignment = Alignment(vertical="top")

        # PLZ als Text, damit Excel nichts umformatiert
        sheet.cell(row=line, column=7).number_format = "@"

        if row["tel"]:
            cell = sheet.cell(row=line, column=3)
            cell.hyperlink = "tel:" + row["tel"].replace(" ", "")
            cell.font = LINK_FONT
        if row["mail"]:
            cell = sheet.cell(row=line, column=4)
            cell.hyperlink = "mailto:" + row["mail"]
            cell.font = LINK_FONT
        if row["web"]:
            cell = sheet.cell(row=line, column=5)
            cell.hyperlink = row["web"]
            cell.font = LINK_FONT
        if row["maps"]:
            cell = sheet.cell(row=line, column=11)
            cell.hyperlink = row["maps"]
            cell.font = LINK_FONT

        sheet.cell(row=line, column=6).fill = PatternFill("solid", fgColor=colour)

    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(rows) + 1}"


def write_overview(sheet, rows):
    sheet.column_dimensions["A"].width = 30
    for column in "BCDE":
        sheet.column_dimensions[column].width = 16
    sheet.column_dimensions["F"].width = 40

    title = sheet.cell(row=1, column=1, value="Schweizer Sportvereine – Kontaktliste")
    title.font = Font(name=ARIAL, size=16, bold=True, color=INK)
    sheet.cell(
        row=2,
        column=1,
        value=f"Stand {STAND}. Golf, Tennis, Reiten und Eishockey aus OpenStreetMap; "
        "Schwimmen aus der Mitgliedvereinsliste von Swiss Aquatics.",
    ).font = Font(name=ARIAL, size=10, italic=True)

    head = ["Sportart", "Einträge", "mit Telefon", "mit E-Mail", "Website ok", "Priorität"]
    for index, value in enumerate(head, start=1):
        cell = sheet.cell(row=4, column=index, value=value)
        cell.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL

    sports = [(sport, prio) for _, sport, prio in SOURCES]
    for offset, (sport, prio) in enumerate(sports):
        line = 5 + offset
        sheet.cell(row=line, column=1, value=sport).font = BODY
        sheet.cell(row=line, column=2, value=f"=COUNTIF('Alle Vereine'!$A:$A,$A{line})").font = BODY
        sheet.cell(
            row=line,
            column=3,
            value=f'=COUNTIFS(\'Alle Vereine\'!$A:$A,$A{line},\'Alle Vereine\'!$C:$C,"<>")',
        ).font = BODY
        sheet.cell(
            row=line,
            column=4,
            value=f'=COUNTIFS(\'Alle Vereine\'!$A:$A,$A{line},\'Alle Vereine\'!$D:$D,"<>")',
        ).font = BODY
        sheet.cell(
            row=line,
            column=5,
            value=f'=COUNTIFS(\'Alle Vereine\'!$A:$A,$A{line},\'Alle Vereine\'!$F:$F,"Link ok*")',
        ).font = BODY
        sheet.cell(row=line, column=6, value=prio).font = BODY

    total = 5 + len(sports)
    sheet.cell(row=total, column=1, value="Total").font = Font(name=ARIAL, size=10, bold=True)
    for column in range(2, 6):
        letter = get_column_letter(column)
        cell = sheet.cell(row=total, column=column, value=f"=SUM({letter}5:{letter}{total - 1})")
        cell.font = Font(name=ARIAL, size=10, bold=True)

    notes = [
        "",
        "Spalte «Link geprüft»",
        "Link ok – Website wurde am " + STAND + " angesteuert und hat geantwortet (HTTP 200/301).",
        "Prüfen – Server blockt automatische Aufrufe (403/406) oder antwortete nicht rechtzeitig; im Browser meist normal erreichbar.",
        "Link tot – Adresse liefert 404 oder einen Serverfehler.",
        "",
        "Was diese Liste NICHT ist",
        "Telefonnummern und Mailadressen stammen 1:1 aus der Quelle und wurden nicht nachtelefoniert.",
        "OpenStreetMap ist nicht vollständig: Vereine ohne eigene Anlage fehlen dort häufig ganz.",
        "Bei Schwimmen publiziert Swiss Aquatics keine Telefonnummern – dort gibt es Website und Ort.",
        "Für eine lückenlose Liste pro Sportart braucht es die Verbandsverzeichnisse (Swiss Tennis, Swiss Golf, football.ch, stv-fsg.ch, sihf.ch).",
        "",
        "Die Spalte «Priorität» kommt aus der Marktanalyse (coachgrid/marktanalyse-vereine.md):",
        "A = eigene Anlage und angestelltes Personal, echtes Budget. B = eigene Anlage, kleines Budget.",
    ]
    for offset, text in enumerate(notes):
        line = total + 2 + offset
        cell = sheet.cell(row=line, column=1, value=text)
        bold = text in ("Spalte «Link geprüft»", "Was diese Liste NICHT ist")
        cell.font = Font(name=ARIAL, size=10, bold=bold)


def write_sources(sheet):
    sheet.column_dimensions["A"].width = 40
    sheet.column_dimensions["B"].width = 78
    sheet.cell(row=1, column=1, value="Quelle").font = Font(
        name=ARIAL, size=10, bold=True, color="FFFFFF"
    )
    sheet.cell(row=1, column=2, value="Link").font = Font(
        name=ARIAL, size=10, bold=True, color="FFFFFF"
    )
    sheet.cell(row=1, column=1).fill = HEAD_FILL
    sheet.cell(row=1, column=2).fill = HEAD_FILL

    entries = [
        ("Vereine Golf/Tennis/Reiten/Eishockey", "https://overpass.osm.ch/api/interpreter"),
        ("OSM-Tag club=sport", "https://wiki.openstreetmap.org/wiki/Tag:club=sport"),
        (
            "Schwimmvereine (Swiss Aquatics)",
            "https://www.swiss-aquatics.ch/verband/mitglieder/alle-mitgliedvereine/",
        ),
        (
            "Mitgliederzahlen pro Schwimmverein",
            "https://www.swiss-aquatics.ch/wp-content/uploads/2024/04/Uebersichtsliste_gemeldete_Mitglieder_2023_2024.pdf",
        ),
        (
            "Vereinsstudie 2022 (Swiss Olympic)",
            "https://www.swissolympic.ch/dam/jcr:e13bfb8d-92a6-41a3-89e4-b80d30c2d23c/Vereinsstudie%202022_DE_Web.pdf",
        ),
        (
            "Vereinsfinanzen (Sportobservatorium)",
            "https://www.sportobs.ch/inhalte/Indikatoren_PDF_neu/Ind_22_Sportobs.pdf",
        ),
        (
            "Sport Schweiz 2020 (BASPO)",
            "https://www.sportobs.ch/inhalte/Downloads/Bro_Sport_Schweiz_2020_d_WEB.pdf",
        ),
        ("Swiss Tennis", "https://www.swisstennis.ch/de/"),
        ("Swiss Golf", "https://www.swissgolf.ch/"),
        ("Swiss Ice Hockey", "https://www.sihf.ch/"),
        ("Schweizerischer Turnverband", "https://www.stv-fsg.ch/"),
    ]
    for offset, (label, url) in enumerate(entries):
        line = offset + 2
        sheet.cell(row=line, column=1, value=label).font = BODY
        cell = sheet.cell(row=line, column=2, value=url)
        cell.hyperlink = url
        cell.font = LINK_FONT

    line = len(entries) + 3
    sheet.cell(
        row=line,
        column=1,
        value="Lizenz: OpenStreetMap-Daten stehen unter ODbL 1.0 – bei Nutzung auf der Website "
        "muss «© OpenStreetMap-Mitwirkende» genannt werden.",
    ).font = Font(name=ARIAL, size=10, italic=True)


def main():
    rows = load()
    book = Workbook()

    overview = book.active
    overview.title = "Übersicht"
    liste = book.create_sheet("Alle Vereine")
    quellen = book.create_sheet("Quellen")

    write_list(liste, rows)
    write_overview(overview, rows)
    write_sources(quellen)

    book.save(OUT)
    print(f"{len(rows)} Einträge -> {OUT}")


if __name__ == "__main__":
    main()
